from pathlib import Path
from io import BytesIO

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.services.profile_service import ProfileService
from app.services.resume_service import ResumeService


PROJECT_ROOT = Path(__file__).resolve().parents[3]
EXPORT_DIR = PROJECT_ROOT / "data" / "exports"


class ResumeExportService:
    def __init__(self) -> None:
        self.profile_service = ProfileService()
        self.resume_service = ResumeService()

    def export_markdown_file(self) -> Path:
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        markdown = self.resume_service.generate_markdown().content
        output_path = EXPORT_DIR / "resume-export.md"
        output_path.write_text(markdown, encoding="utf-8")
        return output_path

    def _build_basic_info_lines(self, profile) -> list[str]:
        """Return (key_label: value) lines for all non-empty basic info fields."""
        bi = profile.basic_info
        cf = bi.custom_fields or []
        fields = [
            ("求职意向", bi.job_intention),
            ("邮箱", bi.email),
            ("电话", bi.phone),
            ("出生年月", bi.birth_date),
            ("最高学历", bi.highest_degree),
        ]
        for f in cf:
            if f.key and f.value:
                fields.append((f.key, f.value))
        return [f"{k}: {v}" for k, v in fields if v]

    def export_template_docx(self) -> Path:
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        profile = self.profile_service.get_profile()
        bi = profile.basic_info

        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.15

        # Title
        title = doc.add_heading(bi.name or "个人简历", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Basic info on one line
        info_parts = self._build_basic_info_lines(profile)
        if info_parts:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(" | ".join(info_parts))
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

        # Helper to add sections
        def add_section(title_text: str):
            doc.add_heading(title_text, level=2)

        # Self intro
        if profile.self_intro:
            add_section("自我简介")
            doc.add_paragraph(profile.self_intro)

        # Skills
        if profile.skills:
            add_section("个人技能")
            doc.add_paragraph("、".join(profile.skills))

        # Work experience
        if profile.work_experiences:
            add_section("工作经历")
            for e in profile.work_experiences:
                p = doc.add_paragraph()
                run = p.add_run(f"{e.company_name} - {e.position_name}")
                run.bold = True
                run.font.size = Pt(11)
                if e.work_time:
                    run2 = p.add_run(f"    {e.work_time}")
                    run2.font.size = Pt(9)
                    run2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                for c in e.work_contents:
                    doc.add_paragraph(c, style="List Bullet")

        # Project experience
        if profile.project_experiences:
            add_section("项目经验")
            for pr in profile.project_experiences:
                p = doc.add_paragraph()
                title_run = p.add_run(pr.project_name or "项目")
                title_run.bold = True
                title_run.font.size = Pt(11)
                if pr.project_role:
                    role_run = p.add_run(f" ({pr.project_role})")
                    role_run.font.size = Pt(10)
                    role_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                if pr.project_time:
                    time_run = p.add_run(f"    {pr.project_time}")
                    time_run.font.size = Pt(9)
                    time_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                if pr.project_description:
                    for line in pr.project_description.split("\n"):
                        if line.strip():
                            doc.add_paragraph(line.strip())

        # Education
        if profile.education:
            add_section("教育背景")
            for e in profile.education:
                p = doc.add_paragraph()
                run = p.add_run(e.school_name or "")
                run.bold = True
                run.font.size = Pt(11)
                if e.school_time:
                    run2 = p.add_run(f"    {e.school_time}")
                    run2.font.size = Pt(9)
                    run2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                if e.major:
                    doc.add_paragraph(f"{e.major}{' | ' + e.degree if e.degree else ''}")

        # Awards
        if profile.awards:
            add_section("获奖情况")
            for a in profile.awards:
                doc.add_paragraph(a, style="List Bullet")

        # Self evaluation
        if profile.self_evaluation:
            add_section("自我评价")
            for e in profile.self_evaluation:
                doc.add_paragraph(e, style="List Bullet")

        output_path = EXPORT_DIR / "resume-export.docx"
        doc.save(str(output_path))
        return output_path

    def export_html(self) -> Path:
        from pathlib import Path
        profile = self.profile_service.get_profile()
        bi = profile.basic_info
        cf = bi.custom_fields or []

        def sec(title, body):
            return '<div class="preview-section"><h3 class="preview-section-title">' + title + '</h3><div class="preview-section-body">' + body + '</div></div>'

        skills_html = "".join('<span class="tag">' + s + '</span>' for s in profile.skills) if profile.skills else '<p class="preview-empty">None</p>'

        work_html = ""
        for e in profile.work_experiences:
            contents = "".join("<li>" + c + "</li>" for c in e.work_contents)
            work_html += '<div class="preview-entry"><div class="preview-entry-title"><span>' + e.company_name + ' - ' + e.position_name + '</span> <span class="preview-entry-time">' + e.work_time + '</span></div><ul class="preview-entry-ul">' + contents + '</ul></div>'
        if not work_html: work_html = '<p class="preview-empty">None</p>'

        proj_html = ""
        for p in profile.project_experiences:
            desc = (p.project_description or "").replace("\n", "<br/>")
            proj_html += '<div class="preview-entry"><div class="preview-entry-title"><span>' + p.project_name + ' (' + (p.project_role or "") + ')</span> <span class="preview-entry-time">' + (p.project_time or "") + '</span></div><p class="preview-text" style="white-space:pre-wrap">' + desc + '</p></div>'
        if not proj_html: proj_html = '<p class="preview-empty">None</p>'

        edu_html = ""
        for e in profile.education:
            degree_span = '<span class="edu-degree"> | ' + e.degree + '</span>' if e.degree else ""
            edu_html += '<div class="preview-entry edu-entry"><div class="preview-entry-title edu-title-line"><span>' + e.school_name + '</span><span class="preview-entry-time">' + e.school_time + '</span></div><div class="edu-detail"><span class="edu-major">' + e.major + '</span>' + degree_span + '</div></div>'
        if not edu_html: edu_html = '<p class="preview-empty">None</p>'

        awards_html = "".join("<li>" + a + "</li>" for a in profile.awards) if profile.awards else '<p class="preview-empty">None</p>'
        self_eval_html = "".join("<li>" + ev + "</li>" for ev in profile.self_evaluation) if profile.self_evaluation else '<p class="preview-empty">None</p>'

        all_fields = [
            ("求职意向", bi.job_intention),
            ("邮箱", bi.email),
            ("电话", bi.phone),
            ("出生年月", bi.birth_date),
            ("最高学历", bi.highest_degree),
        ] + [(f.key, f.value) for f in cf if f.key and f.value]
        meta_fields = "".join('<span class="preview-meta-field">' + k + ': ' + v + '</span>' for k, v in all_fields if v)
        photo_html = '<img src="' + bi.photo + '" class="preview-photo preview-photo-right" />' if bi.photo else ""

        css = "*{margin:0;padding:0;box-sizing:border-box}@page{size:A4;margin:15mm}body{font-family:Microsoft YaHei,sans-serif;background:#fff;padding:0}.preview-doc{width:210mm;min-height:297mm;margin:0 auto;background:#fff;padding:25mm 20mm;border-radius:0;box-shadow:none}@media print{.preview-doc{page-break-after:always}}.preview-header{display:flex;gap:16px;align-items:flex-start;margin-bottom:24px;padding-bottom:16px;border-bottom:2px solid #e8e4f0}.preview-photo{width:80px;height:100px;object-fit:cover;border-radius:4px;flex-shrink:0}.preview-photo-right{margin-left:auto}.preview-basic-left{flex:1}.preview-name{font-size:22px;font-weight:700;color:#1a1a2e;margin-bottom:4px}.preview-meta-fields{display:flex;flex-wrap:wrap;gap:4px 0}.preview-meta-field{white-space:nowrap;font-weight:400;color:#6c6c8a;font-size:12px}.preview-meta-field::after{content:'|';color:#c0bcd0;margin:0 4px}.preview-meta-field:last-child::after{content:none}.preview-section{margin-bottom:18px}.preview-section-title{font-size:11px;font-weight:700;text-transform:uppercase;letter-spacing:.06em;color:#7c3aed;margin-bottom:10px;padding-bottom:4px;border-bottom:1px solid rgba(167,139,250,.2)}.preview-section-body{font-size:13px;color:#3a3a5c;line-height:1.65}.preview-text{font-size:13px;line-height:1.7;color:#3a3a5c}.preview-entry{margin-bottom:12px}.preview-entry-title{font-size:13px;font-weight:600;color:#1a1a2e;margin-bottom:4px;display:flex;justify-content:space-between;align-items:center}.preview-entry-time{font-size:11px;font-weight:400;color:#a0a0b8;flex-shrink:0}.preview-entry-ul{padding-left:16px;margin:4px 0}.preview-entry-ul li{font-size:12px;color:#4a4a6c;margin-bottom:3px}.preview-empty{font-size:12px;color:#c0bcd0;font-style:italic}.tag{padding:4px 10px;border-radius:999px;background:rgba(167,139,250,.08);color:#7c3aed;font-size:12px;font-weight:500;display:inline-block;margin:2px 4px 2px 0}.edu-title-line{display:flex;justify-content:space-between}.edu-detail{font-size:11px;color:#8a8aa0;margin-top:2px}.edu-major{font-size:12px;color:#6c6c8a}.edu-degree{font-size:11px;color:#a0a0b8}@media print{body{background:#fff;padding:0}.preview-doc{box-shadow:none}}"

        html = '<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8"><title>Resume</title><style>' + css + '</style></head><body><div class="preview-doc"><div class="preview-header"><div class="preview-basic-left"><h1 class="preview-name">' + (bi.name or "姓名") + '</h1><div class="preview-meta-fields">' + meta_fields + '</div></div>' + photo_html + '</div>' + sec("自我介绍", '<p class="preview-text">' + (profile.self_intro or "暂无") + '</p>') + sec("个人技能", skills_html) + sec("工作经历", work_html) + sec("项目经验", proj_html) + sec("教育背景", edu_html) + sec("获奖情况", '<ul class="preview-entry-ul">' + awards_html + '</ul>') + sec("自我评价", '<ul class="preview-entry-ul">' + self_eval_html + '</ul>') + '</div></body></html>'

        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        output_path = EXPORT_DIR / "resume-preview.html"
        output_path.write_text(html, encoding="utf-8")
        return output_path



    def export_pdf(self) -> Path:
        return self.export_html()
